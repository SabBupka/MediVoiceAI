from datetime import date, datetime
from pathlib import Path
import sqlite3


class SlotManager:
    """Локальный источник доступных слотов и блокировок."""

    def __init__(self, db_path: str) -> None:
        Path(db_path).parent.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(db_path, check_same_thread=False)
        self.db.execute(
            """CREATE TABLE IF NOT EXISTS slots (
              doctor_id TEXT, doctor_name TEXT, specialty TEXT,
              scheduled_at TEXT, blocked INTEGER DEFAULT 0,
            PRIMARY KEY (doctor_id, scheduled_at)
            )"""
        )
        self.db.execute(
            """CREATE TABLE IF NOT EXISTS bookings (
              visit_id INTEGER PRIMARY KEY, doctor_id TEXT, scheduled_at TEXT
            )"""
        )
        self.db.commit()

    def import_excel(self, path: str, doctor_ids: dict[str, int] | None = None) -> int:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Слоты"]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            return 0
        inserted = 0
        for row in rows[1:]:
            if not row or not row[0] or not row[3] or row[11] != "Свободен":
                continue
            excel_doctor_id, name, specialty = str(row[3]), str(row[4]), str(row[5])
            doctor_id = str((doctor_ids or {}).get(name, excel_doctor_id))
            slot_date = row[1]
            start = row[8]
            if not hasattr(slot_date, "strftime") or not hasattr(start, "strftime"):
                continue
            scheduled_at = f"{slot_date.strftime('%d.%m.%Y')} {start.strftime('%H:%M')}"
            cursor = self.db.execute(
                "INSERT OR IGNORE INTO slots VALUES (?, ?, ?, ?, 0)",
                (doctor_id, name, specialty, scheduled_at),
            )
            inserted += cursor.rowcount
        self.db.commit()
        return inserted

    def list_slots(self, specialty: str | None = None, limit: int = 20) -> list[dict]:
        rows = self.db.execute(
            "SELECT doctor_id, doctor_name, specialty, scheduled_at, blocked FROM slots"
        ).fetchall()
        now = datetime.now()
        groups: dict[tuple[str, str, str], list[tuple]] = {}
        for row in rows:
            doctor_id, name, doctor_specialty, scheduled_at, blocked = row
            if specialty and doctor_specialty.casefold() != specialty.casefold():
                continue
            key = (name.strip().casefold(), doctor_specialty.strip().casefold(), scheduled_at)
            groups.setdefault(key, []).append(row)
        result = []
        for group in groups.values():
            # ID aliases for one doctor/time are a single logical appointment slot.
            # If any alias is blocked, do not leak another duplicate as available.
            if any(row[4] for row in group):
                continue
            doctor_id, name, doctor_specialty, scheduled_at, _ = group[0]
            parsed = datetime.strptime(scheduled_at, "%d.%m.%Y %H:%M")
            if parsed >= now:
                result.append(
                    {
                        "doctor_id": doctor_id,
                        "doctor_name": name,
                        "specialty": doctor_specialty,
                        "scheduled_at": scheduled_at,
                        "_sort": parsed,
                    }
                )
        result.sort(key=lambda item: item["_sort"])
        for item in result:
            item.pop("_sort", None)
        return result[:limit]

    def is_available(self, doctor_id: str, scheduled_at: str) -> bool:
        return self._available_doctor_id(doctor_id, scheduled_at) is not None

    def reserve(
        self,
        doctor_id: str,
        scheduled_at: str,
        visit_id: int,
        *,
        allow_existing_visit: bool = False,
    ) -> None:
        stored = self._available_doctor_id(doctor_id, scheduled_at)
        if stored is None:
            if allow_existing_visit:
                booking = self.db.execute(
                    "SELECT 1 FROM bookings WHERE visit_id=? AND scheduled_at=?",
                    (visit_id, scheduled_at),
                ).fetchone()
                if booking:
                    return
            raise ValueError("Слот отсутствует или уже занят")
        aliases = self._aliases(stored, scheduled_at)
        placeholders = ",".join("?" for _ in aliases)
        cursor = self.db.execute(
            f"UPDATE slots SET blocked=1 WHERE doctor_id IN ({placeholders}) "
            "AND scheduled_at=? AND blocked=0",
            (*aliases, scheduled_at),
        )
        if cursor.rowcount < 1:
            raise ValueError("Слот отсутствует или уже занят")
        self.db.execute(
            "INSERT OR REPLACE INTO bookings(visit_id, doctor_id, scheduled_at) VALUES (?, ?, ?)",
            (visit_id, stored, scheduled_at),
        )
        self.db.commit()

    def _aliases(self, doctor_id: str, scheduled_at: str) -> list[str]:
        row = self.db.execute(
            "SELECT doctor_name, specialty FROM slots "
            "WHERE doctor_id=? AND scheduled_at=? LIMIT 1",
            (doctor_id, scheduled_at),
        ).fetchone()
        if not row:
            return [doctor_id]
        aliases = self.db.execute(
            "SELECT doctor_id FROM slots WHERE doctor_name=? AND specialty=? AND scheduled_at=?",
            (row[0], row[1], scheduled_at),
        ).fetchall()
        return [item[0] for item in aliases] or [doctor_id]

    def _available_doctor_id(self, doctor_id: str, scheduled_at: str) -> str | None:
        candidates = [str(doctor_id)]
        if str(doctor_id).isdigit():
            candidates.append(f"D{int(doctor_id):03d}")
        found = self.db.execute(
            f"SELECT doctor_id, doctor_name, blocked FROM slots "
            f"WHERE doctor_id IN ({','.join('?' for _ in candidates)}) AND scheduled_at=?",
            (*candidates, scheduled_at),
        ).fetchall()
        if not found:
            return None
        logical_rows = self.db.execute(
            "SELECT doctor_id, blocked FROM slots "
            "WHERE doctor_name=? AND specialty=(SELECT specialty FROM slots "
            "WHERE doctor_id=? AND scheduled_at=? LIMIT 1) AND scheduled_at=?",
            (found[0][1], found[0][0], scheduled_at, scheduled_at),
        ).fetchall()
        if any(row[1] for row in logical_rows):
            return None
        return found[0][0]

    def _stored_doctor_id(self, doctor_id: str, scheduled_at: str) -> str | None:
        candidates = [str(doctor_id)]
        if str(doctor_id).isdigit():
            candidates.append(f"D{int(doctor_id):03d}")
        for candidate in candidates:
            if self.db.execute(
                "SELECT 1 FROM slots WHERE doctor_id=? AND scheduled_at=?",
                (candidate, scheduled_at),
            ).fetchone():
                return candidate
        return None

    def release(self, visit_id: int) -> None:
        row = self.db.execute(
            "SELECT doctor_id, scheduled_at FROM bookings WHERE visit_id=?", (visit_id,)
        ).fetchone()
        if not row:
            return
        aliases = self._aliases(row[0], row[1])
        placeholders = ",".join("?" for _ in aliases)
        self.db.execute(
            f"UPDATE slots SET blocked=0 WHERE doctor_id IN ({placeholders}) AND scheduled_at=?",
            (*aliases, row[1]),
        )
        self.db.execute("DELETE FROM bookings WHERE visit_id=?", (visit_id,))
        self.db.commit()

    def release_slot(self, doctor_id: str, scheduled_at: str) -> None:
        stored = self._stored_doctor_id(doctor_id, scheduled_at)
        if not stored:
            return
        aliases = self._aliases(stored, scheduled_at)
        placeholders = ",".join("?" for _ in aliases)
        self.db.execute(
            f"UPDATE slots SET blocked=0 WHERE doctor_id IN ({placeholders}) AND scheduled_at=?",
            (*aliases, scheduled_at),
        )
        self.db.execute(
            "DELETE FROM bookings WHERE doctor_id IN (" + placeholders + ") AND scheduled_at=?",
            (*aliases, scheduled_at),
        )
        self.db.commit()

    def finalize_reservation(self, temporary_visit_id: int, visit_id: int) -> None:
        row = self.db.execute(
            "SELECT doctor_id, scheduled_at FROM bookings WHERE visit_id=?",
            (temporary_visit_id,),
        ).fetchone()
        if not row:
            raise ValueError("Временная блокировка слота не найдена")
        self.db.execute(
            "DELETE FROM bookings WHERE visit_id=?", (visit_id,)
        )
        self.db.execute(
            "UPDATE bookings SET visit_id=? WHERE visit_id=?",
            (visit_id, temporary_visit_id),
        )
        self.db.commit()
